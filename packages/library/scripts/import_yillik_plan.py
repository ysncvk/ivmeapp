#!/usr/bin/env python3
"""
Yıllık plan Excel dosyasını (fenbilimleri.org biçimi) ayrıştırır ve JSON üretir.

Kullanım:
  python3 import_yillik_plan.py <plan.xlsx> <çıktı.ts> [--prefix FB] [--unit-prefix unite-]
                                [--year 2026]
Gereksinim: openpyxl (yalnızca geliştirme zamanı aracı, uygulamanın bağımlılığı değil).

Çıktı: weeks, breaks, examPeriods dışa aktaran bir TypeScript dosyası. Tarihler ISO
(YYYY-MM-DD); Eylül-Aralık için ilk yıl, Ocak-Haziran için ikinci yıl sayılır. Çıktıyı
`biome` ile biçimlendirin.
"""
import json
import re
import sys
from datetime import date

import openpyxl

MONTHS = {
    "ocak": 1, "şubat": 2, "mart": 3, "nisan": 4, "mayıs": 5, "haziran": 6,
    "temmuz": 7, "ağustos": 8, "eylül": 9, "ekim": 10, "kasım": 11, "aralık": 12,
}


def tr_lower(text: str) -> str:
    return text.replace("İ", "i").replace("I", "ı").lower()


def year_of(month: int, first_year: int) -> int:
    return first_year if month >= 8 else first_year + 1


def parse_range(text: str, first_year: int):
    """'14–18 Eylül' veya '28 Eylül–2 Ekim' -> (başlangıç, bitiş) ISO."""
    t = tr_lower(text.replace("\n", " "))
    m = re.search(
        r"(\d+)\s*([a-zçğıöşü]+)?\s*[–-]\s*(\d+)\s*([a-zçğıöşü]+)", t
    )
    if not m:
        return None
    d1, mon1, d2, mon2 = m.groups()
    m2 = MONTHS[mon2]
    m1 = MONTHS[mon1] if mon1 else m2
    start = date(year_of(m1, first_year), m1, int(d1))
    end = date(year_of(m2, first_year), m2, int(d2))
    return start.isoformat(), end.isoformat()


def parse_period(text: str, first_year: int):
    """'28 Aralık 2026–8 Ocak 2027' veya '2–13 Kasım 2026' (yıllar isteğe bağlı)."""
    t = tr_lower(text.replace("\n", " "))
    m = re.search(
        r"(\d+)\s*([a-zçğıöşü]+)?\s*(20\d\d)?\s*[–-]\s*(\d+)\s*([a-zçğıöşü]+)\s*(20\d\d)?", t
    )
    if not m:
        return None
    d1, mon1, y1, d2, mon2, y2 = m.groups()
    m2 = MONTHS[mon2]
    m1 = MONTHS[mon1] if mon1 else m2
    end_year = int(y2) if y2 else year_of(m2, first_year)
    start_year = int(y1) if y1 else (end_year if not mon1 or m1 == m2 else year_of(m1, first_year))
    return date(start_year, m1, int(d1)).isoformat(), date(end_year, m2, int(d2)).isoformat()


def clean(v) -> str:
    return "" if v is None else re.sub(r"[ \t]+", " ", str(v)).strip()


def main():
    src, out = sys.argv[1], sys.argv[2]

    def opt(name, default):
        return sys.argv[sys.argv.index(name) + 1] if name in sys.argv else default

    prefix = opt("--prefix", "FB")
    unit_prefix = opt("--unit-prefix", "unite-")
    first_year = int(opt("--year", "2026"))
    ws = openpyxl.load_workbook(src, data_only=True).worksheets[0]
    weeks, breaks, exams = [], [], []
    for r in range(6, ws.max_row + 1):
        a, b, c, d, e, f, g, h, i = (clean(ws.cell(r, k).value) for k in range(1, 10))
        if a and not b and ("TATİL" in a or "Bayramı" in a):
            # Hücrede satır sonu olabilir: tek satıra indirip ilk tarih aralığını al
            flat = a.replace("\n", " / ")
            m = re.search(r"(?:TATİLİ|TATİL|Bayramı):\s*([^•/]+?)(?:\s+20\d\d)?\s*(?:[•/]|$)", flat)
            span = parse_range(m.group(1), first_year) if m else None
            title = flat.split(":")[0].strip(" /")
            if span:
                breaks.append({"title": title, "start": span[0], "end": span[1]})
            continue
        mw = re.match(r"(\d+)\.\s*Hafta", b)
        if not mw:
            continue
        # Tarih kısmı 'N. Hafta: \n 14–18 Eylül' biçiminde
        span = parse_range(re.sub(r"^\d+\.\s*Hafta:?", "", b), first_year)
        # Ünite numaraları (başlık satır sonlarıyla bölünmüş olabilir)
        unit_ids = [
            f"{unit_prefix}{n}"
            for n in sorted({int(x) for x in re.findall(r"(\d+)\.\s*ÜNİTE", d)})
        ]
        # Konular: satır sonlarıyla ayrılmış; ham metni satır satır alıyoruz
        raw_e = str(ws.cell(r, 5).value or "")
        topics = [clean(x) for x in raw_e.split("\n") if clean(x)]
        # Kazanımlar: kodlara göre bölünür, satır sonları boşluğa çevrilir
        raw_f = re.sub(r"\s+", " ", str(ws.cell(r, 6).value or "")).strip()
        outcomes = []
        for om in re.finditer(
            rf"({prefix}\.\d+\.\d+\.\d+\.\d+)\.\s*(.*?)(?=\s{prefix}\.\d+\.\d+\.\d+\.\d+\.|$)", raw_f
        ):
            outcomes.append({"code": om.group(1), "title": om.group(2).strip()})
        otp = re.search(r"(\d+)\s*saati\s*Okul Temelli", i)
        weeks.append(
            {
                "week": int(mw.group(1)),
                "start": span[0],
                "end": span[1],
                "hours": c,
                "unitIds": unit_ids,
                "topics": topics,
                "outcomes": outcomes,
                **({"otpHours": int(otp.group(1))} if otp else {}),
                **({"note": re.sub(r"\s+", " ", i).strip()} if i else {}),
            }
        )
    for r in range(6, ws.max_row + 1):
        h = str(ws.cell(r, 8).value or "")
        em = re.search(r"SINAV HAFTALARI\s*\(([^)]+)\)", h)
        if em:
            span = parse_period(em.group(1), first_year)
            if span:
                exams.append({"title": "Sınav haftaları", "start": span[0], "end": span[1]})
    dump = lambda v: json.dumps(v, ensure_ascii=False, indent=1)
    with open(out, "w", encoding="utf-8") as fh:
        fh.write(
            "// Otomatik üretildi: packages/library/scripts/import_yillik_plan.py. Elle düzenlemeyin.\n"
            'import type { CurriculumPeriod, CurriculumWeek } from "@ivmeapp/core";\n\n'
            f"export const weeks: CurriculumWeek[] = {dump(weeks)};\n\n"
            f"export const breaks: CurriculumPeriod[] = {dump(breaks)};\n\n"
            f"export const examPeriods: CurriculumPeriod[] = {dump(exams)};\n"
        )
    print(f"{len(weeks)} hafta, {len(breaks)} tatil, {len(exams)} sınav dönemi")


if __name__ == "__main__":
    main()
